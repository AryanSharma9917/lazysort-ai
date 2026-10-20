"""Create a disposable, fictional folder for demonstrating LazySort."""
from pathlib import Path

from docx import Document
from pypdf import PdfWriter


ROOT = Path(__file__).resolve().parents[1] / "demo_workspace"


def main() -> None:
    ROOT.mkdir(exist_ok=True)
    samples = {
        "meeting-notes.txt": "Project Atlas meeting, 20 October 2026. Next action: prepare launch checklist.",
        "project-plan.md": "# Project Atlas\nMilestone: fictional beta release. Owner: Alex Example.",
        "expenses.csv": "date,vendor,amount\n2026-10-12,Example Hosting,24.00\n",
        "resume-draft.txt": "Taylor Example — software engineer. Fictional demonstration document.",
        "random-notes.txt": "Ideas for the fictional community garden: seeds, signs, watering rota.",
        "data.json": '{"project":"Atlas","status":"planning","fictional":true}',
    }
    for name, content in samples.items():
        (ROOT / name).write_text(content, encoding="utf-8")
    document = Document()
    document.add_heading("Fictional Quarterly Report", 0)
    document.add_paragraph("Project Atlas completed its fictional planning milestone.")
    document.save(ROOT / "document-final.docx")
    PdfWriter().write(ROOT / "invoice_oct.pdf")
    print(f"Created safe demo files in {ROOT}")


if __name__ == "__main__":
    main()
